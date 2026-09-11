#!/usr/bin/env python3
"""Generate DOCX materials for software copyright application.

Standard requirements (中国版权保护中心):
- 程序鉴别材料: 60页（连续的前30页 + 后30页），纯代码，不含说明性文字
- 文档鉴别材料: 60页（连续的前30页 + 后30页），软件设计/使用说明文档
- 程序每页50个非空有效行，文档每页30个非空有效行
- 正文不添加行号；不足时禁止使用空白或循环模板填充
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    print("Error: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

sys.path.insert(0, str(Path(__file__).parent))

from material_requirements import (
    assert_generation_allowed,
    assert_output_path_allowed,
    load_material_requirements,
)
from validate_materials import _substantive_duplicates, audit_pdf


# 标准参数统一来自 material-requirements.yaml
MATERIAL_REQUIREMENTS = load_material_requirements()
PAGE_SELECTION = MATERIAL_REQUIREMENTS['page_selection']
PROGRAM_POLICY = MATERIAL_REQUIREMENTS['program_material']
DOCUMENT_POLICY = MATERIAL_REQUIREMENTS['document_material']
LINES_PER_PAGE = int(PROGRAM_POLICY['lines_per_page'])
DOC_LINES_PER_PAGE = int(DOCUMENT_POLICY['lines_per_page'])
FRONT_PAGES = int(PAGE_SELECTION['front_pages'])
BACK_PAGES = int(PAGE_SELECTION['back_pages'])
TOTAL_PAGES = int(PAGE_SELECTION['total_pages'])
PROGRAM_FONT_SIZE = float(PROGRAM_POLICY['font']['generation_point_size'])
DOCUMENT_FONT_SIZE = float(DOCUMENT_POLICY['font']['generation_point_size'])


def _assert_generation_target(
    manifest: dict,
    output_dir: Path,
    software_name: str,
    requirements: dict | None = None,
) -> None:
    rules = requirements or MATERIAL_REQUIREMENTS
    assert_generation_allowed(
        str(manifest.get('copyrightId', '')),
        software_name,
        rules,
    )
    assert_output_path_allowed(output_dir, rules)


def render_and_verify_pdf(
    docx_path: Path,
    material_type: str,
    requirements: dict | None = None,
) -> Path:
    """Render DOCX to PDF and verify the final pagination, rows, font and line numbers."""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        raise RuntimeError('未找到 LibreOffice/soffice，无法完成最终PDF渲染与页数复核')
    with tempfile.TemporaryDirectory(prefix='copyright-soffice-') as profile:
        completed = subprocess.run(
            [
                soffice,
                f'-env:UserInstallation={Path(profile).resolve().as_uri()}',
                '--headless',
                '--convert-to',
                'pdf',
                '--outdir',
                str(docx_path.parent),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
    pdf_path = docx_path.with_suffix('.pdf')
    if completed.returncode != 0 or not pdf_path.is_file():
        details = (completed.stderr or completed.stdout).strip()
        raise RuntimeError(f'DOCX转PDF失败：{details or "未生成PDF文件"}')
    report = audit_pdf(
        pdf_path,
        material_type,
        requirements or MATERIAL_REQUIREMENTS,
    )
    if not report['passed']:
        messages = '；'.join(issue['message'] for issue in report['issues'][:10])
        raise ValueError(f'渲染后的{pdf_path.name}未通过硬性审计：{messages}')
    print(
        f"Verified: {pdf_path} ({report['page_count']} pages, "
        f"minimum {report['minimum_body_rows']} rows/page, "
        f"minimum {report['minimum_body_font_size']} pt)"
    )
    return pdf_path


def _add_footer_page_number(section) -> None:
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run('第 ')
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, end])
    run.add_text(f' 页 / 共 {TOTAL_PAGES} 页')
    run.font.size = Pt(8)


def create_code_document(title: str, version: str) -> Document:
    """Create a new DOCX document for code material (程序鉴别材料)."""
    doc = Document()
    
    # 设置页面边距（标准：上2.5cm 下2.5cm 左3cm 右2.5cm）
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
        _add_footer_page_number(section)
    
    # 添加页眉（软件名称 + 版本号）
    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = f"{title} {version}"
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for run in header_para.runs:
        run.font.size = Pt(9)
    
    return doc


def create_doc_document(title: str, version: str) -> Document:
    """Create a new DOCX document for document material (文档鉴别材料)."""
    doc = Document()
    
    # 设置页面边距
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
        _add_footer_page_number(section)
    
    # 页眉
    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = f"{title} {version} 设计说明书"
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for run in header_para.runs:
        run.font.size = Pt(9)
    
    return doc


def read_source_files(workspace_dir: Path, manifest: dict) -> list:
    """Read all source files from workspace.
    
    Returns raw code content only - no added annotations or explanatory text.
    """
    source_dir = workspace_dir / 'source'
    files_content = []
    
    for file_info in manifest['allowedFiles']:
        file_path = source_dir / file_info['path']
        
        if not file_path.exists():
            continue
        
        try:
            content = file_path.read_text(encoding='utf-8', errors='ignore')
            lines = content.split('\n')
            
            # 严格口径：空白行不得计入每页有效行数。
            effective_lines = [line.rstrip() for line in lines if line.strip()]
            
            files_content.append({
                'path': file_info['path'],
                'language': file_info['language'],
                'lines': lines,
                'effective_lines': effective_lines,
                'total_lines': len(lines),
                'effective_line_count': len(effective_lines),
            })
        except Exception as e:
            print(f"Warning: Could not read {file_path}: {e}")
    
    return files_content


def collect_pure_code_lines(files_content: list) -> list:
    """Collect all code lines from files - PURE CODE ONLY.
    
    No file path annotations, no separator comments, no explanatory text.
    Only the actual source code content from the repository.
    """
    all_lines = []
    
    for file_info in files_content:
        # 直接添加代码行，不添加任何说明性文字
        for line in file_info['effective_lines']:
            all_lines.extend(_wrap_print_line(line, int(PROGRAM_POLICY['max_print_columns'])))
    
    return all_lines


def _char_width(char: str) -> int:
    return 2 if unicodedata.east_asian_width(char) in 'WFA' else 1


def _wrap_print_line(text: str, max_columns: int) -> list[str]:
    """Wrap a non-empty print row without adding line numbers or filler."""
    if not text.strip():
        return []
    rows = []
    current = []
    width = 0
    for char in text.rstrip():
        char_width = _char_width(char)
        if current and width + char_width > max_columns:
            rows.append(''.join(current).rstrip())
            current = [char]
            width = char_width
        else:
            current.append(char)
            width += char_width
    if current:
        rows.append(''.join(current).rstrip())
    return [row for row in rows if row.strip()]


def paginate_lines(all_lines: list, lines_per_page: int = LINES_PER_PAGE) -> list:
    """Split effective rows into fixed-size pages."""
    pages = []
    for i in range(0, len(all_lines), lines_per_page):
        page_lines = all_lines[i:i + lines_per_page]
        pages.append(page_lines)
    return pages


def select_material_rows(all_rows: list[str], lines_per_page: int) -> list[str]:
    """Select complete first-30 and last-30 page rows without padding."""
    required_rows = TOTAL_PAGES * lines_per_page
    if len(all_rows) < required_rows:
        raise ValueError(
            f'有效材料只有{len(all_rows)}行，至少需要{required_rows}行；'
            '请补充与系统功能相关的真实内容，禁止空白或循环模板填充'
        )
    if len(all_rows) == required_rows:
        return all_rows
    front_rows = FRONT_PAGES * lines_per_page
    back_rows = BACK_PAGES * lines_per_page
    return all_rows[:front_rows] + all_rows[-back_rows:]


def write_code_to_docx(doc: Document, pages: list, is_last_section: bool = True):
    """Write code pages to DOCX document.
    
    Pure code output - monospace font, no decorations.
    """
    # 设置默认字体为等宽字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Courier New'
    font.size = Pt(PROGRAM_FONT_SIZE)
    
    for page_idx, page_lines in enumerate(pages):
        for line_idx, line in enumerate(page_lines):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            
            # 添加代码文本
            if not line.strip():
                raise ValueError('程序鉴别材料不得写入空白占位行')
            run = para.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(PROGRAM_FONT_SIZE)
        
        # 每页后添加分页符（除了最后一页）
        is_last_page = (page_idx == len(pages) - 1) and is_last_section
        if not is_last_page:
            doc.add_page_break()


def generate_code_material(
    workspace_dir: Path,
    output_dir: Path,
    software_name: str,
    version: str,
    requirements: dict | None = None,
):
    """Generate program identification material (程序鉴别材料).
    
    Standard: 60 pages (front 30 + back 30), pure code only.
    No explanatory text, no file path annotations - only repository code.
    """
    # 加载 manifest
    manifest_path = workspace_dir / 'SOURCE_MANIFEST.json'
    with open(manifest_path, 'r', encoding='utf-8') as f:
        manifest = json.load(f)
    rules = requirements or MATERIAL_REQUIREMENTS
    _assert_generation_target(manifest, output_dir, software_name, rules)
    
    # 读取源文件
    files_content = read_source_files(workspace_dir, manifest)
    
    if not files_content:
        raise ValueError('工作区中没有可用于程序鉴别材料的真实源文件')
    
    # 收集纯代码行（不含任何说明性文字）
    all_lines = collect_pure_code_lines(files_content)
    total_lines = len(all_lines)
    
    # 分页
    pages = paginate_lines(all_lines)
    total_pages = len(pages)
    
    print(f"Total code lines: {total_lines}")
    print(f"Total pages: {total_pages} ({LINES_PER_PAGE} lines/page)")
    
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # 按有效行选取连续前30页和连续后30页，尾页不以空白补齐。
    selected_rows = select_material_rows(all_lines, LINES_PER_PAGE)
    selected_pages = paginate_lines(selected_rows, LINES_PER_PAGE)
    if len(selected_pages) != TOTAL_PAGES:
        raise AssertionError(f'程序鉴别材料输出页数不是{TOTAL_PAGES}页')
    if any(len(page) != LINES_PER_PAGE or any(not line.strip() for line in page)
           for page in selected_pages):
        raise AssertionError('程序鉴别材料必须逐页包含50个非空有效行')

    print(
        f"Generating strict {TOTAL_PAGES}-page material "
        f"(front {FRONT_PAGES} + back {BACK_PAGES}, no line numbers)..."
    )
    doc = create_code_document(software_name, version)
    write_code_to_docx(doc, selected_pages, is_last_section=True)
    output_path = output_dir / f"{software_name}-程序鉴别材料.docx"
    doc.save(str(output_path))
    print(f"Generated: {output_path}")
    print(f"  Pages: {len(selected_pages)}, rows/page: {LINES_PER_PAGE}")
    render_and_verify_pdf(output_path, 'program', rules)
    
    # 生成代码清单（辅助文件，不属于鉴别材料）
    manifest_output = output_dir / f"{software_name}-源代码清单.json"
    with open(manifest_output, 'w', encoding='utf-8') as f:
        json.dump({
            'software_name': software_name,
            'version': version,
            'total_files': len(files_content),
            'total_effective_lines': total_lines,
            'total_pages': total_pages,
            'output_pages': TOTAL_PAGES,
            'front_pages': FRONT_PAGES,
            'back_pages': BACK_PAGES,
            'lines_per_page': LINES_PER_PAGE,
            'line_numbers': False,
            'files': [
                {
                    'path': f['path'],
                    'language': f['language'],
                    'total_lines': f['total_lines'],
                    'effective_lines': f['effective_line_count'],
                }
                for f in files_content
            ]
        }, f, indent=2, ensure_ascii=False)
    
    print(f"Generated: {manifest_output}")
    return output_path


def generate_doc_material(
    workspace_dir: Path,
    output_dir: Path,
    software_name: str,
    version: str,
    description: str = "",
    manual_path: Path | None = None,
    requirements: dict | None = None,
):
    """Generate document identification material (文档鉴别材料).
    
    Standard: 60 pages (front 30 + back 30).
    Content: Software design document / user manual structure.
    
    The body must come from a system-specific manual supplied by the applicant.
    Generic descriptions are never fabricated to reach the page threshold.
    """
    # 加载 manifest
    manifest_path = workspace_dir / 'SOURCE_MANIFEST.json'
    with open(manifest_path, 'r', encoding='utf-8') as f:
        manifest = json.load(f)
    rules = requirements or MATERIAL_REQUIREMENTS
    _assert_generation_target(manifest, output_dir, software_name, rules)
    
    # 读取源文件（用于分析结构）
    files_content = read_source_files(workspace_dir, manifest)
    
    if not files_content:
        raise ValueError('工作区中没有与说明书对应的真实源文件')
    
    output_dir.mkdir(parents=True, exist_ok=True)
    
    copyright_id = str(manifest.get('copyrightId', workspace_dir.name))
    manual_candidates = [
        manual_path,
        workspace_dir / 'SOFTWARE_MANUAL.md',
        workspace_dir.parents[1] / 'manuals' / f'{copyright_id}.md',
    ]
    selected_manual = next(
        (Path(path) for path in manual_candidates if path and Path(path).is_file()),
        None,
    )
    if selected_manual is None:
        raise FileNotFoundError(
            '缺少系统专用说明书。请通过 --manual 指定 Markdown 文件，或放置到 '
            f'{workspace_dir / "SOFTWARE_MANUAL.md"} / '
            f'{workspace_dir.parents[1] / "manuals" / f"{copyright_id}.md"}'
        )
    manual_text = selected_manual.read_text(encoding='utf-8', errors='ignore')
    _validate_manual_source(selected_manual, manual_text)
    if '![' in manual_text or '<img' in manual_text.lower():
        raise ValueError('含图片的说明书请使用 generate_markdown_pdf.py 的 Markdown -> PDF 主流程')
    doc_lines = manual_text.splitlines()

    # 合并已验证的材料元数据与人工核实的系统专用说明书。
    doc_content = _build_document_content(
        software_name, version, description, files_content, doc_lines
    )
    
    # 严格正文行：去除空白、控制打印宽度、拒绝重复和占位内容。
    doc_content = _normalize_document_rows(doc_content)
    _validate_document_content(doc_content)

    # 文档鉴别材料按30个非空正文行分页。
    pages = paginate_lines(doc_content, DOC_LINES_PER_PAGE)
    total_pages = len(pages)
    
    print(f"Document material: {len(doc_content)} lines, {total_pages} pages")
    
    # 按有效行选取连续前30页和连续后30页，确保每页均为完整30行。
    selected_rows = select_material_rows(doc_content, DOC_LINES_PER_PAGE)
    selected_pages = paginate_lines(selected_rows, DOC_LINES_PER_PAGE)
    if len(selected_pages) != TOTAL_PAGES:
        raise AssertionError(f'文档鉴别材料输出页数不是{TOTAL_PAGES}页')
    if any(len(page) != DOC_LINES_PER_PAGE or any(not line.strip() for line in page)
           for page in selected_pages):
        raise AssertionError('文档鉴别材料必须逐页包含30个非空有效行')

    doc = create_doc_document(software_name, version)
    write_doc_pages_to_docx(doc, selected_pages, is_last_section=True)
    
    output_path = output_dir / f"{software_name}-文档鉴别材料.docx"
    doc.save(str(output_path))
    print(f"Generated: {output_path}")
    print(f"  Pages: {len(selected_pages)}, rows/page: {DOC_LINES_PER_PAGE}")
    render_and_verify_pdf(output_path, 'document', rules)
    
    return output_path


def _normalize_document_rows(lines: list[str]) -> list[str]:
    rows: list[str] = []
    max_columns = int(DOCUMENT_POLICY['max_print_columns'])
    for line in lines:
        stripped = line.strip() if line else ''
        if not stripped:
            continue
        if stripped.startswith('```') or re.fullmatch(r'[-*_]{3,}', stripped):
            continue
        if re.fullmatch(r'\|?(?:\s*:?-{3,}:?\s*\|)+', stripped):
            continue
        if re.fullmatch(r'<!--.*?-->', stripped):
            continue
        rows.extend(_wrap_print_line(stripped, max_columns))
    return rows


def _validate_document_content(lines: list[str]) -> None:
    required_lines = TOTAL_PAGES * DOC_LINES_PER_PAGE
    if len(lines) < required_lines:
        raise ValueError(
            f'文档只有{len(lines)}个有效正文行，至少需要{required_lines}行；'
            '请补充与系统功能、操作、异常处理和验收相关的真实内容，'
            '禁止空白或循环模板填充'
        )

    forbidden = MATERIAL_REQUIREMENTS['generation']['forbidden_placeholder_markers']
    joined = '\n'.join(lines).lower()
    for marker in forbidden:
        if marker.lower() in joined:
            raise ValueError(f'文档包含禁止的占位内容：{marker}')

    normalized = [
        re.sub(r'\d+(?:\.\d+)*', '<N>', re.sub(r'\s+', '', line))
        for line in lines
        if len(line.strip()) >= 30
    ]
    duplicates = [line for line, count in Counter(normalized).items() if count > 1]
    if duplicates:
        raise ValueError(
            f'文档存在{len(duplicates)}条重复的实质内容，禁止循环填充；'
            f'首个重复内容：{duplicates[0][:60]}'
        )


def _validate_manual_source(path: Path, text: str) -> None:
    """Validate the authored manual before it is converted into fixed print rows."""
    forbidden = MATERIAL_REQUIREMENTS['generation']['forbidden_placeholder_markers']
    for marker in forbidden:
        if marker.lower() in text.lower():
            raise ValueError(f'说明书{path}包含禁止的占位内容：{marker}')

    headings = re.findall(r'^#{1,6}\s+\S.+$', text, flags=re.MULTILINE)
    numbered_steps = re.findall(
        r'^\s*(?:\d+(?:\.\d+)*[.、]?|[（(]\d+[）)])\s*\S+',
        text,
        flags=re.MULTILINE,
    )
    if len(headings) < 3 or not numbered_steps:
        raise ValueError('说明书必须使用清晰的章节层级、分段和操作步骤编号')

    tips_rule = DOCUMENT_POLICY['supplemental_operation_tips']
    tips_match = re.search(
        rf'^#{{1,6}}\s+.*{re.escape(tips_rule["heading"])}.*$',
        text,
        flags=re.MULTILINE,
    )
    if not tips_match:
        raise ValueError(f'说明书缺少“{tips_rule["heading"]}”章节')
    tips_and_after = text[tips_match.end():]
    if not re.search(r'^#{2,6}\s+\S+', tips_and_after, flags=re.MULTILINE) and not re.search(
        r'^\s*(?:\d+(?:\.\d+)*[.、]?|[（(]\d+[）)])\s*\S+',
        tips_and_after,
        flags=re.MULTILINE,
    ):
        raise ValueError('“补充操作提示”之后必须分段并使用小节或步骤编号')

    duplicates = _substantive_duplicates(tips_and_after)
    if duplicates:
        raise ValueError(
            f'“补充操作提示”之后存在{len(duplicates)}个重复内容块，禁止循环填充；'
            f'首个重复内容：{duplicates[0][:80]}'
        )


def _build_document_content(
    software_name: str,
    version: str,
    description: str,
    files_content: list,
    manual_lines: list,
) -> list[str]:
    """Combine verified metadata with an authored, system-specific manual."""
    main_function_rule = DOCUMENT_POLICY['main_function_section']
    main_function_characters = len(re.sub(r'\s+', '', description or ''))
    minimum = int(main_function_rule['minimum_characters'])
    maximum = int(main_function_rule['maximum_characters'])
    if not minimum <= main_function_characters <= maximum:
        raise ValueError(
            f'《软件的主要功能》必须为{minimum}至{maximum}字；'
            f'当前description为{main_function_characters}字'
        )
    if not manual_lines:
        raise ValueError('系统专用说明书不得为空')

    languages = sorted({item['language'] for item in files_content if item['language']})
    return [
        f'# {software_name}软件说明书',
        '## 一、申请材料信息',
        f'**软件名称：** {software_name}',
        f'**版本号：** {version}',
        '### 1.1 软件的主要功能',
        description.strip(),
        '### 1.2 材料来源说明',
        f'本说明书对应{len(files_content)}个已锁定源文件，主要开发语言为'
        f'{"、".join(languages) if languages else "以源文件清单为准"}。',
        '## 二、软件说明书正文',
        *manual_lines,
    ]


def write_doc_pages_to_docx(doc: Document, pages: list, is_last_section: bool = True):
    """Write document pages to DOCX (文档鉴别材料 uses Song font)."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'SimSun'
    font.size = Pt(DOCUMENT_FONT_SIZE)
    
    for page_idx, page_lines in enumerate(pages):
        for line in page_lines:
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            
            if not line.strip():
                raise ValueError('文档鉴别材料不得写入空白占位行')
            heading_match = re.match(r'^(#{1,6})\s+(.+)$', line.strip())
            display_text = heading_match.group(2) if heading_match else line
            display_text = re.sub(r'^\*\*(.+?)\*\*$', r'\1', display_text)
            run = para.add_run(display_text)
            run.font.name = 'SimSun'
            run.font.size = Pt(DOCUMENT_FONT_SIZE)
            run.bold = bool(heading_match)
        
        is_last_page = (page_idx == len(pages) - 1) and is_last_section
        if not is_last_page:
            doc.add_page_break()


def main():
    """Route the legacy entry point to the Markdown -> PDF workflow."""
    print('generate_docx.py 已转入兼容模式，主流程为 Markdown -> PDF -> 复核')
    from generate_markdown_pdf import main as markdown_pdf_main

    markdown_pdf_main()


if __name__ == '__main__':
    main()
